from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt


def export_presentation(
    insights: list[str],
    summary: str,
    output_path: str = "executive_summary.pptx",
) -> str:
    if len(insights) != 5:
        raise ValueError("Exactly five insights are required.")

    prs = Presentation()
    prs.core_properties.title = "Data Insights Executive Summary"
    prs.core_properties.subject = "W8D1 - Groq Executive Summary"

    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Executive Data Summary"
    slide.placeholders[1].text = "W8D1 | Data Analyst Intern"

    # Slides 2–6: Five insights
    for index, insight in enumerate(insights, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Insight {index}"
        body = slide.placeholders[1]
        body.text = insight
        for paragraph in body.text_frame.paragraphs:
            paragraph.font.size = Pt(24)

    # Slide 7: Summary and recommendations
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Summary & Recommendations"
    body = slide.placeholders[1]
    body.text = summary + "\n\nRecommended next steps:"
    recommendations = [
        "Validate key findings with business stakeholders.",
        "Investigate the drivers behind the observed patterns.",
        "Track relevant KPIs and review changes regularly.",
    ]
    for item in recommendations:
        paragraph = body.text_frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(18)

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return str(path.resolve())
