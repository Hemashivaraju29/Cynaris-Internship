import sys
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from business_recommendations import (
    calculate_insights,
    generate_business_narrative,
)
from ppt_export import export_presentation


@pytest.fixture
def sales_df():
    return pd.DataFrame({
        "Region": ["North", "South", "South", "North"],
        "Sales": [1000, 2500, 1500, 800],
        "Category": ["A", "B", "A", "B"],
    })


def test_calculate_returns_five_insights(sales_df):
    insights = calculate_insights(sales_df)
    assert len(insights) == 5
    assert "4 rows and 3 columns" in insights[0]
    assert "South" in insights[4]


def test_empty_dataframe_raises_error():
    with pytest.raises(ValueError, match="empty"):
        calculate_insights(pd.DataFrame())


def test_dataset_without_numeric_column_raises_error():
    df = pd.DataFrame({"Region": ["North", "South"]})
    with pytest.raises(ValueError, match="numeric"):
        calculate_insights(df)


def test_narrative_requires_five_insights():
    with pytest.raises(ValueError, match="five insights"):
        generate_business_narrative(["Only one insight"])


def test_export_creates_seven_slides(tmp_path):
    insights = [f"Insight {i}" for i in range(1, 6)]
    narrative = (
        "Executive Summary\nSales performance overview.\n"
        "Business Recommendations\n1. Review sales trends."
    )
    output = tmp_path / "recommendations.pptx"

    saved_path = export_presentation(insights, narrative, str(output))

    assert Path(saved_path).exists()
    presentation = Presentation(saved_path)
    assert len(presentation.slides) == 7
    assert "Business Recommendations" in presentation.slides[6].shapes.title.text


def test_export_rejects_wrong_number_of_insights(tmp_path):
    with pytest.raises(ValueError, match="five insights"):
        export_presentation(
            ["Only one insight"],
            "Some recommendations",
            str(tmp_path / "invalid.pptx"),
        )


def test_export_rejects_empty_narrative(tmp_path):
    with pytest.raises(ValueError, match="cannot be empty"):
        export_presentation(
            [f"Insight {i}" for i in range(1, 6)],
            " ",
            str(tmp_path / "invalid.pptx"),
        )
