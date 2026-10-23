from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def export_presentation(
    insights: list[str],
    narrative: str,
   output_path: str = "Week 08/Day 3/output/business_recommendations.pptx",
) -> str:
    """Export five data insights and business narrative to a 7-slide deck."""
    if len(insights) != 5:
        raise ValueError("Exactly five insights are required.")

    if not narrative or not narrative.strip():
        raise ValueError("The business narrative cannot be empty.")

    presentation = Presentation()
    presentation.core_properties.title = "Business Recommendations from Data Insights"
    presentation.core_properties.subject = "W8D3 - PPT Report Export"
    presentation.core_properties.author = "Data Analyst Intern"

    # Slide 1: Title
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Business Insights & Recommendations"
    slide.placeholders[1].text = "W8D3 | Executive Data Report"

    # Slides 2–6: Five insights
    for index, insight in enumerate(insights, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Data Insight {index}"

        body = slide.placeholders[1]
        body.text = insight
        body.text_frame.word_wrap = True

        for paragraph in body.text_frame.paragraphs:
            paragraph.font.size = Pt(24)

    # Slide 7: Executive summary and business recommendations
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Executive Summary & Business Recommendations"

    text_box = slide.shapes.add_textbox(
        Inches(0.6),
        Inches(1.3),
        Inches(8.8),
        Inches(5.7),
    )
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(0.15)
    text_frame.margin_right = Inches(0.15)

    # Split the narrative into paragraphs for readability.
    paragraphs = [part.strip() for part in narrative.split("\n") if part.strip()]

    for index, content in enumerate(paragraphs):
        paragraph = (
            text_frame.paragraphs[0]
            if index == 0
            else text_frame.add_paragraph()
        )
        paragraph.text = content
        paragraph.font.size = Pt(14)
        paragraph.space_after = Pt(8)

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)

    return str(path.resolve())
